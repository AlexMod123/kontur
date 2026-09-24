"""КОНТУР-ПРО: экспорт в Excel (опциональная зависимость openpyxl)."""

from __future__ import annotations

import importlib.util
import io
from typing import TYPE_CHECKING, Any

from kontur.config import CATEGORIES, PRICE_LIST, VERSION, VERSION_NAME
from kontur.export.base import ExportContext, ExportFormat
from kontur.export.registry import register

if TYPE_CHECKING:
    from openpyxl import Workbook


def _bold_header(ws: Any, headers: list[str]) -> None:
    """Записывает и выделяет жирным первую (заголовочную) строку листа."""
    from openpyxl.styles import Font

    for col, h in enumerate(headers, 1):
        ws.cell(1, col, h).font = Font(bold=True)


def _build_summary_sheet(wb: Workbook, ctx: ExportContext) -> None:
    r = ctx.result
    ws = wb.active
    ws.title = "Сводка"
    _bold_header(ws, ["Параметр", "Значение"])
    rows = [
        ("Версия", f"v{VERSION} «{VERSION_NAME}»"),
        ("Дата", ctx.display_timestamp),
        ("Установленная мощность, кВт", round(r.installed_power_w / 1000, 2)),
        ("Расчётная мощность, кВт", round(r.demand_power_w / 1000, 2)),
        ("Ток, А", r.total_current_a),
        ("cos φ", r.cos_phi),
        ("Перекос фаз, %", r.phase_imbalance_pct),
        ("Автомат", f"С{r.recommended_breaker}"),
        ("Кабель", r.recommended_cable),
        ("Падение напряжения, %", r.voltage_drop_pct),
        ("Заземление, Ом", r.grounding_r),
        ("Охлаждение, кВт", round(r.cooling_required_w / 1000, 1)),
        ("ИБП, кВт", round(r.ups_power_w / 1000, 1)),
        ("Токи КЗ 3ф, А", r.short_circuit_3ph),
        ("Токи КЗ 1ф, А", r.short_circuit_1ph),
        ("Утечки, мА", r.leakage_current_ma),
        ("УЗО, мА", r.leakage_uzo_ma),
        ("Молниезащита", r.lightning_zone),
        ("ЛВС, портов", r.lan_ports),
        ("ЛВС, кабель м", r.lan_cable_m),
        ("PoE, Вт (треб.)", r.poe_required_w),
        ("PoE, Вт (бюджет)", r.poe_budget_w),
        ("Теплопотери, Вт", r.heat_loss_w),
        ("Шум, дБ", r.noise_level_db),
        ("Смета, ₽", r.cost_total),
    ]
    for row, (k, v) in enumerate(rows, 2):
        ws.cell(row, 1, k)
        ws.cell(row, 2, v)


def _build_devices_sheet(wb: Workbook, ctx: ExportContext) -> None:
    ws = wb.create_sheet("Оборудование")
    _bold_header(
        ws, ["№", "Обозначение", "Наименование", "Категория", "Мощность, Вт", "Напряжение", "Цена, ₽"]
    )
    for row, d in enumerate(ctx.devices, 2):
        ws.cell(row, 1, row - 1)
        ws.cell(row, 2, d.designation or "—")
        ws.cell(row, 3, d.name)
        ws.cell(row, 4, CATEGORIES.get(d.category, {}).get("name", d.category))
        ws.cell(row, 5, d.power_w)
        ws.cell(row, 6, d.voltage)
        ws.cell(row, 7, d.price or PRICE_LIST.get(d.equip_type, 0))


def _build_rooms_sheet(wb: Workbook, ctx: ExportContext) -> None:
    ws = wb.create_sheet("Помещения")
    _bold_header(ws, ["№", "Название", "Площадь, м²", "ПК", "Камеры", "СКУД", "ОПС"])
    for row, room in enumerate(ctx.rooms, 2):
        ws.cell(row, 1, row - 1)
        ws.cell(row, 2, room.name)
        ws.cell(row, 3, room.area)
        ws.cell(row, 4, room.pc_count)
        ws.cell(row, 5, room.camera_count)
        ws.cell(row, 6, "Да" if room.has_skud else "—")
        ws.cell(row, 7, "Да" if room.has_ops else "—")


def _build_checks_sheet(wb: Workbook, ctx: ExportContext) -> None:
    if not ctx.checks:
        return
    ws = wb.create_sheet("Проверки")
    _bold_header(ws, ["Проверка", "Статус", "Детали"])
    for row, c in enumerate(ctx.checks, 2):
        ws.cell(row, 1, c["name"])
        ws.cell(row, 2, c["status"])
        ws.cell(row, 3, c["detail"])


def _build_spec_sheet(wb: Workbook, ctx: ExportContext) -> None:
    if not ctx.spec:
        return
    ws = wb.create_sheet("Спецификация")
    _bold_header(ws, ["№", "Обозначение", "Наименование", "Тип", "Вендор", "Ед.", "Кол.", "Примечание"])
    for row, s in enumerate(ctx.spec, 2):
        ws.cell(row, 1, s.get("pos", ""))
        ws.cell(row, 2, s.get("designation", ""))
        ws.cell(row, 3, s.get("name", ""))
        ws.cell(row, 4, s.get("type", ""))
        ws.cell(row, 5, s.get("vendor", ""))
        ws.cell(row, 6, s.get("unit", ""))
        ws.cell(row, 7, s.get("qty", ""))
        ws.cell(row, 8, s.get("note", ""))


def _build_cable_journal_sheet(wb: Workbook, ctx: ExportContext) -> None:
    if not ctx.cable_journal:
        return
    ws = wb.create_sheet("Кабельный журнал")
    _bold_header(ws, ["№", "Обозначение", "Начало", "Конец", "Кабель", "Сечение", "Длина, м", "Ток, А"])
    for row, c in enumerate(ctx.cable_journal, 2):
        ws.cell(row, 1, c.get("num", ""))
        ws.cell(row, 2, c.get("designation", ""))
        ws.cell(row, 3, c.get("start", ""))
        ws.cell(row, 4, c.get("end", ""))
        ws.cell(row, 5, c.get("cable", ""))
        ws.cell(row, 6, c.get("section", ""))
        ws.cell(row, 7, c.get("length_m", ""))
        ws.cell(row, 8, c.get("current_a", ""))


def _build_recommendations_sheet(wb: Workbook, ctx: ExportContext) -> None:
    ws = wb.create_sheet("Рекомендации")
    _bold_header(ws, ["№", "Текст"])
    for row, rec in enumerate(ctx.result.recommendations, 2):
        ws.cell(row, 1, row - 1)
        ws.cell(row, 2, rec)


def _build_phase_balance_sheet(wb: Workbook, ctx: ExportContext) -> None:
    ws = wb.create_sheet("Баланс фаз")
    _bold_header(ws, ["Фаза", "Ток, А"])
    for row, (phase, current) in enumerate(ctx.result.per_phase_current.items(), 2):
        ws.cell(row, 1, phase)
        ws.cell(row, 2, current)
    ws.cell(5, 1, "Перекос, %")
    ws.cell(5, 2, ctx.result.phase_imbalance_pct)


@register
class ExcelExportFormat(ExportFormat):
    """Экспорт полного отчёта в книгу Excel (8 листов). Требует openpyxl."""

    name = "excel"
    extension = "xlsx"

    @staticmethod
    def is_available() -> bool:
        return importlib.util.find_spec("openpyxl") is not None

    @staticmethod
    def render(ctx: ExportContext) -> bytes:
        from openpyxl import Workbook

        wb = Workbook()
        _build_summary_sheet(wb, ctx)
        _build_devices_sheet(wb, ctx)
        _build_rooms_sheet(wb, ctx)
        _build_checks_sheet(wb, ctx)
        _build_spec_sheet(wb, ctx)
        _build_cable_journal_sheet(wb, ctx)
        _build_recommendations_sheet(wb, ctx)
        _build_phase_balance_sheet(wb, ctx)

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
